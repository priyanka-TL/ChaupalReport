"""
3_final_processor_v2.py — Optimized Final Report Processor
============================================================
Key optimizations vs original (3_final_processor.py):
  1. REFINE_BATCH_SIZE increased 30→60 (~50% fewer refinement API calls)
  2. Unified semantic dedup replaces per-theme + cross-theme AI clustering
     (~20 calls → 2-4 calls)
  3. Batched insight generation: 5 concepts per API call
     (~50 calls → ~10 calls)
  4. Cross-theme consolidation simplified to Phase 1 only (exact-match, 0 API calls)
  5. Cross-batch dedup simplified (larger batch, fewer sub-batches)

Estimated total API calls: ~20 (was ~80) — 75% reduction
Same output structure, accuracy, and report format.
"""

import pandas as pd
import re
import os
import time
import random
from difflib import SequenceMatcher
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import json
from dotenv import load_dotenv
from llm_provider import LLMProvider

load_dotenv()

# --- LLM CONFIGURATION ---
try:
    llm_provider = LLMProvider()
except Exception as e:
    print(f"⚠️ LLM Client Setup Failed: {e}")
    llm_provider = None

THEME_KNOWLEDGE_BASE = """
1. Poverty and Economic Barriers
2. Legal Document-linked Barriers
3. Child Marriage Issue
4. Distance and Accessibility Issues
5. Parental Attitudes & Socio-Cultural
6. School Infrastructure & Facility
7. Teacher Capacity & Quality
8. Safety Issues
9. Substance Abuse & Addiction
10. Other Factors
"""

THEME_NAMES = [
    "Poverty And Economic Barriers",
    "Legal Document-Linked Barriers",
    "Child Marriage Issue",
    "Distance And Accessibility Issues",
    "Parental Attitudes & Socio-Cultural",
    "School Infrastructure & Facility",
    "Teacher Capacity & Quality",
    "Safety Issues",
    "Substance Abuse & Addiction",
    "Other Factors",
]

MAX_RETRIES = int(os.getenv("LLM_MAX_RETRIES", "5"))
BASE_RETRY_SECONDS = float(os.getenv("LLM_RETRY_BASE_SECONDS", "2"))
MAX_RETRY_SECONDS = float(os.getenv("LLM_RETRY_MAX_SECONDS", "45"))
REFINE_CHECKPOINT_DIR = os.getenv("REFINE_CHECKPOINT_DIR", ".")
# ══ OPTIMIZATION 1: Larger refine batch (was 30) ══
REFINE_BATCH_SIZE = int(os.getenv("REFINE_BATCH_SIZE", "60"))
REFINE_MAX_TOKENS = int(os.getenv("REFINE_MAX_TOKENS", "6000"))
REFINE_THINKING_BUDGET = int(os.getenv("REFINE_THINKING_BUDGET", "256"))
INSIGHT_MAX_TOKENS = int(os.getenv("INSIGHT_MAX_TOKENS", "900"))
INSIGHT_THINKING_BUDGET = int(os.getenv("INSIGHT_THINKING_BUDGET", "256"))
INSIGHT_TEMPERATURE = float(os.getenv("INSIGHT_TEMPERATURE", "0.15"))
# ══ OPTIMIZATION 2: Unified dedup batch size ══
UNIFIED_DEDUP_BATCH_SIZE = int(os.getenv("UNIFIED_DEDUP_BATCH_SIZE", "150"))
# ══ OPTIMIZATION 3: Insight batch size (concepts per API call) ══
INSIGHT_BATCH_SIZE = int(os.getenv("INSIGHT_BATCH_SIZE", "5"))
# ══ OPTIMIZATION 5: Larger cross-batch dedup (was 80) ══
CROSS_DEDUP_BATCH = int(os.getenv("CROSS_DEDUP_BATCH", "150"))


# ═══════════════════════════════════════════
# UTILITY FUNCTIONS
# ═══════════════════════════════════════════

def is_retryable_error(error):
    message = str(error).lower()
    retry_signals = [
        "429",
        "rate limit",
        "resource_exhausted",
        "too many requests",
        "throttle",
        "temporarily unavailable",
        "timeout",
        "max_tokens",
        "missing text parts",
        "thinking budget",
    ]
    return any(signal in message for signal in retry_signals)


def get_refinement_checkpoint_path(type_label):
    safe_label = str(type_label).strip().lower().replace(" ", "_")
    return os.path.join(REFINE_CHECKPOINT_DIR, f"refine_checkpoint_{safe_label}.json")


def load_refinement_checkpoint(checkpoint_path):
    if not os.path.exists(checkpoint_path):
        return {}
    try:
        with open(checkpoint_path, "r", encoding="utf-8") as file:
            data = json.load(file)
        return data if isinstance(data, dict) else {}
    except Exception as error:
        print(f"   ⚠️ Could not read refinement checkpoint {checkpoint_path}: {error}")
        return {}


def save_refinement_checkpoint(checkpoint_path, results):
    os.makedirs(os.path.dirname(checkpoint_path) or ".", exist_ok=True)
    with open(checkpoint_path, "w", encoding="utf-8") as file:
        json.dump(results, file, ensure_ascii=False, indent=2)


def _extract_json_block(text):
    cleaned = str(text).replace('```json', '').replace('```', '').strip()
    start = cleaned.find('{')
    end = cleaned.rfind('}')
    if start != -1 and end != -1 and end > start:
        return cleaned[start:end + 1]
    return cleaned


def _parse_refinement_response(text, batch):
    raw_json = _extract_json_block(text)
    parsed = json.loads(raw_json)
    if not isinstance(parsed, dict):
        raise ValueError("Refinement response is not a JSON object")

    theme_name_set = {t.lower() for t in THEME_NAMES}
    number_map = {
        "1": "Poverty And Economic Barriers",
        "2": "Legal Document-Linked Barriers",
        "3": "Child Marriage Issue",
        "4": "Distance And Accessibility Issues",
        "5": "Parental Attitudes & Socio-Cultural",
        "6": "School Infrastructure & Facility",
        "7": "Teacher Capacity & Quality",
        "8": "Safety Issues",
        "9": "Substance Abuse & Addiction",
        "10": "Other Factors",
    }

    valid = {}
    for item in batch:
        value = parsed.get(item)
        if not isinstance(value, dict):
            continue
        concept = str(value.get('concept', '')).strip()
        theme = str(value.get('theme', '')).strip()

        # Map numeric theme codes to full names
        if theme.isdigit():
            theme = number_map.get(theme, theme)

        # Avoid collapsing concepts into theme names
        if concept.lower() in theme_name_set:
            concept = item

        if concept and theme:
            valid[item] = {'concept': concept, 'theme': theme}

    if not valid:
        raise ValueError("No valid refinement mappings found in response")

    return valid


def _repair_refinement_json_with_ai(raw_text, batch, type_label):
    if not llm_provider:
        return None

    repair_prompt = f"""Fix the malformed JSON below.

RULES:
1. Return ONLY valid JSON object.
2. Keys must come from this input list: {json.dumps(batch, ensure_ascii=False)}
3. Each value must be an object with keys: concept, theme.
4. Keep meaning intact. Do not add extra text.

MALFORMED JSON/TEXT:
{raw_text}

OUTPUT: valid JSON object only."""

    repaired = llm_provider.generate_text(repair_prompt, max_tokens=2500, temperature=0)
    return _parse_refinement_response(repaired, batch)


# ═══════════════════════════════════════════
# CLASSIFICATION FUNCTIONS
# ═══════════════════════════════════════════

def categorize_environment_aggressive(text):
    """Ultra-Aggressive Environment Classification to minimize Unmapped tags."""
    text_lower = str(text).lower()

    school_kw = ['school', 'teacher', 'classroom', 'class', 'student', 'education', 'study', 'teaching', 'academic', 'admission', 'enroll', 'attendance', 'grade', 'subject', 'exam', 'books', 'uniform', 'midday meal', 'mid day', 'scholarship', 'library', 'playground', 'infrastructure', 'facility', 'toilet', 'water', 'building']
    home_kw = ['parent', 'family', 'mother', 'father', 'home', 'household', 'house', 'sibling', 'brother', 'sister', 'domestic', 'child labour', 'work at home', 'income', 'alcoholic', 'migration', 'marriage', 'dowry', 'attitude', 'mindset', 'belief', 'cultural', 'discrimination']
    comm_kw = ['village', 'community', 'society', 'road', 'transport', 'bus', 'distance', 'far', 'path', 'route', 'weather', 'rain', 'heat', 'flood', 'surroundings', 'neighborhood', 'area', 'locality', 'safety', 'harassment', 'molestation', 'social pressure', 'caste', 'tribe', 'practice']

    s_score = sum(2 if kw in text_lower else 0 for kw in school_kw)
    h_score = sum(2 if kw in text_lower else 0 for kw in home_kw)
    c_score = sum(2 if kw in text_lower else 0 for kw in comm_kw)

    if any(kw in text_lower for kw in ['to school', 'reach school', 'go to school']): c_score += 3
    if any(kw in text_lower for kw in ['at home', 'in family', 'parent awareness']): h_score += 3
    if any(kw in text_lower for kw in ['in school', 'at school', 'lacks']): s_score += 3

    scores = {'School': s_score, 'Home': h_score, 'Community': c_score}
    if max(scores.values()) == 0:
        if any(w in text_lower for w in ['poor', 'poverty', 'money', 'financial']): return 'Home'
        return 'Community'
    return max(scores, key=scores.get)


def categorize_agency(text):
    """Classifies the driver of the solution."""
    text_lower = str(text).lower()
    comm_kw = ['community', 'together', 'collective', 'meena manch', 'chaupal', 'village', 'we will', 'committee', 'panchayat']
    ind_kw = ['parent', 'family', 'individual', 'we should', 'people should', 'personally', 'mother', 'father']
    inst_kw = ['government', 'school', 'ngo', 'administration', 'authority', 'provide', 'officer', 'department', 'teacher']

    scores = {
        'Community-led': sum(1 for kw in comm_kw if kw in text_lower),
        'Individual-led': sum(1 for kw in ind_kw if kw in text_lower),
        'Institutional': sum(1 for kw in inst_kw if kw in text_lower)
    }
    return max(scores, key=scores.get) if max(scores.values()) > 0 else 'Community-led'


# ═══════════════════════════════════════════
# FORMATTING UTILITIES
# ═══════════════════════════════════════════

def set_cell_background(cell, fill_color):
    shading_elm = OxmlElement('w:shd')
    shading_elm.set(qn('w:fill'), fill_color)
    cell._tc.get_or_add_tcPr().append(shading_elm)


def normalize_text(text):
    if pd.isna(text): return "Uncategorized"
    return re.sub(r'^\d+[\.\)\s-]*', '', str(text)).strip()


def is_valid_solution(text):
    """Filters out problem statements disguised as solutions."""
    text = str(text).lower().strip()
    problem_starters = ['lack of', 'no ', 'not enough', 'poor ', 'insufficient', 'scarcity', 'absence', 'shortage', 'due to', 'because of']
    if any(text.startswith(p) for p in problem_starters):
        return False
    return True


def clean_theme_name(text):
    """Cleans theme names, handling combined themes and empty values."""
    if pd.isna(text) or str(text).strip() == "" or str(text).lower() == "nan":
        return "Other Factors"

    raw = str(text).strip()
    number_map = {
        "1": "Poverty And Economic Barriers",
        "2": "Legal Document-Linked Barriers",
        "3": "Child Marriage Issue",
        "4": "Distance And Accessibility Issues",
        "5": "Parental Attitudes & Socio-Cultural",
        "6": "School Infrastructure & Facility",
        "7": "Teacher Capacity & Quality",
        "8": "Safety Issues",
        "9": "Substance Abuse & Addiction",
        "10": "Other Factors",
    }
    if re.fullmatch(r"\d+[\.\)\s-]*", raw):
        return number_map.get(raw.strip(". )-"), "Other Factors")

    text = raw
    if '+' in text:
        text = text.split('+')[0].strip()
    text = re.sub(r'^\d+[\.\)\s-]*', '', text).strip()
    if not text:
        number_match = re.match(r"^\d+", raw)
        if number_match:
            return number_map.get(number_match.group(0), "Other Factors")
        return "Other Factors"
    text = text.title()

    allowed = {
        "Poverty And Economic Barriers",
        "Legal Document-Linked Barriers",
        "Child Marriage Issue",
        "Distance And Accessibility Issues",
        "Parental Attitudes & Socio-Cultural",
        "School Infrastructure & Facility",
        "Teacher Capacity & Quality",
        "Safety Issues",
        "Substance Abuse & Addiction",
        "Other Factors",
    }
    return text if text in allowed else "Other Factors"


def normalize_concept_key(text):
    """Canonical key for grouping near-duplicate concept labels."""
    if pd.isna(text):
        return ""
    normalized = str(text).strip().lower()
    normalized = re.sub(r'\s+', ' ', normalized)
    normalized = re.sub(r'[^a-z0-9\s]', '', normalized)
    return normalized.strip()


# ═══════════════════════════════════════════
# TOKEN / SIMILARITY FUNCTIONS
# ═══════════════════════════════════════════

def _tokenize_concept_key(text):
    if not text:
        return set()

    stopwords = {
        'a', 'an', 'the', 'and', 'or', 'of', 'to', 'for', 'in', 'on', 'at', 'by',
        'with', 'from', 'is', 'are', 'was', 'were', 'be', 'being', 'this', 'that'
    }

    tokens = []
    for token in str(text).split():
        token = token.strip()
        if not token or token in stopwords:
            continue
        if token.endswith('ation') and len(token) > 7:
            token = token[:-5]
        elif token.endswith('tion') and len(token) > 6:
            token = token[:-4]
        if token.endswith('ing') and len(token) > 5:
            token = token[:-3]
        elif token.endswith('ed') and len(token) > 4:
            token = token[:-2]
        elif token.endswith('es') and len(token) > 4:
            token = token[:-2]
        elif token.endswith('s') and len(token) > 3:
            token = token[:-1]
        if token and token not in stopwords:
            tokens.append(token)

    return set(tokens)


def _soft_token_overlap(tokens_a, tokens_b):
    """Counts token overlap allowing near-lexical matches (generic, no domain keywords)."""
    if not tokens_a or not tokens_b:
        return 0

    used_b = set()
    overlap = 0

    for token_a in tokens_a:
        best_index = None
        best_score = 0.0
        for index, token_b in enumerate(tokens_b):
            if index in used_b:
                continue
            score = SequenceMatcher(None, token_a, token_b).ratio()
            if score > best_score:
                best_score = score
                best_index = index
        if best_index is not None and best_score >= 0.8:
            used_b.add(best_index)
            overlap += 1

    return overlap


def _subject_token(text):
    """Extract the first meaningful stemmed token from the original phrase order."""
    _stopwords_ext = {
        'a', 'an', 'the', 'and', 'or', 'of', 'to', 'for', 'in', 'on', 'at', 'by',
        'with', 'from', 'is', 'are', 'was', 'were', 'be', 'being', 'this', 'that',
        'as', 'due', 'its', 'it',
        'prevent', 'preventing', 'barrier', 'barriers', 'issue', 'issues',
        'cause', 'causing', 'affect', 'affecting', 'hinder', 'hindering',
        'stop', 'stopping', 'impact', 'impacting', 'lead', 'leading',
    }
    for word in str(text).split():
        w = word.strip().lower()
        if not w or w in _stopwords_ext:
            continue
        if w.endswith('ation') and len(w) > 7:
            w = w[:-5]
        elif w.endswith('tion') and len(w) > 6:
            w = w[:-4]
        if w.endswith('ing') and len(w) > 5:
            w = w[:-3]
        elif w.endswith('ed') and len(w) > 4:
            w = w[:-2]
        elif w.endswith('es') and len(w) > 4:
            w = w[:-2]
        elif w.endswith('s') and len(w) > 3:
            w = w[:-1]
        if w and w not in _stopwords_ext:
            return w
    return ""


def _are_concepts_similar(key_a, key_b):
    if not key_a or not key_b:
        return False
    if key_a == key_b:
        return True

    tokens_a = _tokenize_concept_key(key_a)
    tokens_b = _tokenize_concept_key(key_b)
    if not tokens_a or not tokens_b:
        return False

    tokens_a_list = sorted(tokens_a)
    tokens_b_list = sorted(tokens_b)
    soft_intersection = _soft_token_overlap(tokens_a_list, tokens_b_list)

    union = len(tokens_a) + len(tokens_b) - soft_intersection
    jaccard = (soft_intersection / union) if union else 0.0
    overlap = soft_intersection / min(len(tokens_a), len(tokens_b))
    seq_ratio = SequenceMatcher(None, key_a, key_b).ratio()

    subject_sim = 0.0
    if len(tokens_a) >= 3 and len(tokens_b) >= 3:
        subj_a = _subject_token(key_a)
        subj_b = _subject_token(key_b)
        if subj_a and subj_b:
            subject_sim = SequenceMatcher(None, subj_a, subj_b).ratio()
            if subject_sim < 0.70:
                return False

    if subject_sim >= 0.90:
        return (overlap >= 0.45 and jaccard >= 0.25) or seq_ratio >= 0.88

    return (overlap >= 0.6 and jaccard >= 0.40) or seq_ratio >= 0.88


def assign_concept_groups(df, concept_column='Merged_Concept'):
    """Generic concept clustering without hardcoded domain keywords or extra API calls."""
    if df.empty:
        temp = df.copy()
        temp['Concept_Key'] = ""
        temp['Concept_Group'] = ""
        return temp

    temp = df.copy()
    temp['Concept_Key'] = temp[concept_column].apply(normalize_concept_key)
    temp = temp[temp['Concept_Key'] != ""].copy()
    if temp.empty:
        temp['Concept_Group'] = ""
        return temp

    key_counts = temp['Concept_Key'].value_counts().to_dict()
    keys = list(key_counts.keys())

    parent = {key: key for key in keys}

    def find(x):
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    def union(a, b):
        root_a, root_b = find(a), find(b)
        if root_a == root_b:
            return
        if key_counts.get(root_a, 0) >= key_counts.get(root_b, 0):
            parent[root_b] = root_a
        else:
            parent[root_a] = root_b

    for i, key_a in enumerate(keys):
        for key_b in keys[i + 1:]:
            if _are_concepts_similar(key_a, key_b):
                union(key_a, key_b)

    group_map = {key: find(key) for key in keys}
    temp['Concept_Group'] = temp['Concept_Key'].map(group_map)
    return temp


# ═══════════════════════════════════════════
# INSIGHT GENERATION (OPTIMIZED — BATCHED)
# ═══════════════════════════════════════════

def _clean_text_output(text):
    """Clean up text for grammar, spacing, and formatting issues."""
    # Remove multiple spaces
    text = re.sub(r' +', ' ', text)
    # Fix spacing before punctuation
    text = re.sub(r'\s+([.,;:!?])', r'\1', text)
    # Fix spacing after punctuation
    text = re.sub(r'([.,;:!?])([A-Za-z])', r'\1 \2', text)
    # Remove duplicate consecutive words
    text = re.sub(r'\b(\w+)\s+\1\b', r'\1', text, flags=re.IGNORECASE)
    # Ensure sentences end with proper punctuation
    lines = text.split('\n')
    cleaned_lines = []
    for line in lines:
        line = line.strip()
        if line and not line[-1] in '.!?':
            line += '.'
        cleaned_lines.append(line)
    text = '\n'.join(cleaned_lines)
    # Capitalize first letter of sentences
    text = re.sub(r'(^|[.!?]\s+)([a-z])', lambda m: m.group(1) + m.group(2).upper(), text)
    # Fix common run-on issues
    text = re.sub(r'\.([A-Z])', r'. \1', text)
    return text.strip()


def _challenge_item_insight_fallback(theme, concept, share, count, districts, env_text):
    """Generates a deterministic fallback insight (no API call)."""
    return (
        f"This represents {share:.1f}% of theme challenges, indicating a systemic issue rather than isolated incidents. "
        f"The pattern manifests primarily in {env_text} settings, pointing to where interventions must be anchored. "
        f"With {districts} district(s) reporting this challenge, it requires {'localized' if districts <= 2 else 'coordinated multi-district'} response strategies."
    )


def _batch_challenge_insights(theme, concepts_data, t_c, t_s):
    """
    ══ OPTIMIZATION 3: Batched insight generation ══
    Generate insights for multiple challenge concepts in a SINGLE API call.
    Processes INSIGHT_BATCH_SIZE concepts per call instead of 1 per call.

    Args:
        theme: Theme name
        concepts_data: list of dicts with keys: concept, count, share, concept_key, concept_rows, original_texts
        t_c: theme challenge DataFrame
        t_s: theme solution DataFrame

    Returns: dict of {concept_name: [insight_strings]}
    """
    if not llm_provider or not concepts_data:
        return {}

    results = {}
    batch_size = INSIGHT_BATCH_SIZE

    for batch_start in range(0, len(concepts_data), batch_size):
        batch = concepts_data[batch_start:batch_start + batch_size]

        # Build per-concept sections for the prompt
        concept_sections = []
        for idx, data in enumerate(batch, 1):
            concept = data['concept']
            count = data['count']
            share = data['share']
            rows = data.get('concept_rows', pd.DataFrame())
            districts = rows['District'].nunique() if not rows.empty and 'District' in rows.columns else 0

            env_text = "Not available"
            if not rows.empty and 'Environment' in rows.columns:
                env_mix = rows['Environment'].value_counts(normalize=True)
                if not env_mix.empty:
                    env_text = env_mix.index[0]

            texts = data.get('original_texts', [])[:5]  # 5 samples per concept (reduced from 8)
            scenarios = " || ".join(str(t) for t in texts)

            concept_sections.append(
                f"CONCEPT {idx}: {concept}\n"
                f"- Mentions: {count} ({share:.1f}% of theme)\n"
                f"- Geographic spread: {districts} district(s)\n"
                f"- Primary setting: {env_text}\n"
                f"- Ground scenarios: {scenarios}"
            )

        concepts_text = "\n\n".join(concept_sections)

        prompt = f"""You are analyzing on-ground education barriers from grassroots dialogue data.

THEME: {theme}

{concepts_text}

For EACH concept above, write 2-3 DISTINCT, NON-REPETITIVE insights (max 80 words per concept).

EACH INSIGHT MUST COVER A DIFFERENT DIMENSION:
- Insight 1: What MECHANISM/TRIGGER causes this barrier?
- Insight 2: WHO is most affected and WHAT cascading effects occur?
- Insight 3 (if needed): What SYSTEMIC PATTERN or broader implication emerges?

CRITICAL REQUIREMENTS:
- NO repetition — each sentence must add NEW information
- NO generic statements like "barriers impede access" or "factors prevent education"
- NO explicit references to "voices," "testimonials," or "participants said"
- NO repetition of the challenge concept name (it's already in the heading)
- Be CONCRETE and SPECIFIC about mechanisms, triggers, affected groups
- Synthesize the BROADER PICTURE from scenarios — what patterns emerge?
- Use PERFECT grammar, spelling, and punctuation

OUTPUT: Valid JSON object where keys are the EXACT concept names and values are arrays of insight strings.
Example: {{"Concept Name Here": ["First insight sentence.", "Second insight sentence."]}}
RETURN ONLY VALID JSON. NO MARKDOWN. NO EXTRA TEXT."""

        for attempt in range(1, MAX_RETRIES + 1):
            try:
                response = llm_provider.generate_text(
                    prompt,
                    max_tokens=INSIGHT_MAX_TOKENS * len(batch),
                    temperature=INSIGHT_TEMPERATURE,
                    thinking_budget=INSIGHT_THINKING_BUDGET,
                )

                parsed = json.loads(_extract_json_block(response))
                if not isinstance(parsed, dict):
                    raise ValueError("Expected JSON object")

                for concept_data in batch:
                    concept = concept_data['concept']
                    insights = parsed.get(concept, None)

                    # Try exact match first, then fuzzy match
                    if not insights or not isinstance(insights, list):
                        for key, val in parsed.items():
                            if isinstance(val, list) and SequenceMatcher(None, key.lower(), concept.lower()).ratio() > 0.85:
                                insights = val
                                break

                    if insights and isinstance(insights, list):
                        cleaned = [_clean_text_output(str(i)) for i in insights if str(i).strip()]
                        # Enforce per-concept word limit
                        total_text = " ".join(cleaned)
                        words = total_text.split()
                        if len(words) > 100:
                            total_text = " ".join(words[:100]).rstrip(" ,;:") + "."
                            cleaned = [_clean_text_output(total_text)]
                        if cleaned:
                            results[concept] = cleaned

                break  # Success — exit retry loop

            except Exception as e:
                if is_retryable_error(e) and attempt < MAX_RETRIES:
                    delay = min(MAX_RETRY_SECONDS, BASE_RETRY_SECONDS * (2 ** (attempt - 1)))
                    delay += random.uniform(0, 0.5)
                    print(f"      ⚠️ Batch insight attempt {attempt}/{MAX_RETRIES} failed: {e}. Retrying in {delay:.1f}s...")
                    time.sleep(delay)
                    continue
                print(f"      ⚠️ Batch insight generation failed: {e}. Using fallbacks for this batch.")
                break

    return results


def _solution_item_insight(theme, concept, share):
    return (
        f"Insight: This solution reflects a practical response under '{theme}', and contributes "
        f"{share:.1f}% of the proposed actions in this theme."
    )


# ═══════════════════════════════════════════
# AI REFINEMENT & DEDUPLICATION (OPTIMIZED)
# ═══════════════════════════════════════════

THEME_KNOWLEDGE_BASE = """
1. Poverty and Economic Barriers
2. Legal Document-linked Barriers
3. Child Marriage Issue
4. Distance and Accessibility Issues
5. Parental Attitudes & Socio-Cultural
6. School Infrastructure & Facility
7. Teacher Capacity & Quality
8. Safety Issues
9. Substance Abuse & Addiction
10. Other Factors
"""


def refine_concepts_with_ai(concepts_list, type_label):
    """
    Uses AI to clean, deduplicate, and re-theme the top concepts.
    ══ OPTIMIZATION 1: Larger batch size (60 vs 30) — ~50% fewer API calls ══
    """
    if not llm_provider:
        return {}

    checkpoint_path = get_refinement_checkpoint_path(type_label)
    all_results = load_refinement_checkpoint(checkpoint_path)
    pending_concepts = [concept for concept in concepts_list if concept not in all_results]

    print(f"   🧠 AI Refinement: Optimizing top {len(concepts_list)} {type_label}s...")
    if all_results:
        print(f"      ♻️ Resume mode: loaded {len(all_results)} cached refinements from {checkpoint_path}")

    if not pending_concepts:
        print(f"      ✅ No pending refinement for {type_label}. Using cached results.")
        return all_results

    batch_size = REFINE_BATCH_SIZE

    total_batches = (len(pending_concepts) + batch_size - 1) // batch_size
    for i in range(0, len(pending_concepts), batch_size):
        batch = pending_concepts[i:i+batch_size]
        current_batch = (i // batch_size) + 1
        print(f"      Processing batch {current_batch}/{total_batches} ({len(batch)} items)...")

        prompt = f"""You are a Data Cleaning Expert for an Education Report.

THEMES:
{THEME_KNOWLEDGE_BASE}

CONTEXT: This input list contains canonical {type_label} labels generated by an AI across SEPARATE
batches. The SAME concept frequently received DIFFERENT labels in different batches. Your job is
to collapse ALL such variants into one canonical label per concept.

════════════════════════════════════════
UNIVERSAL MERGE PRINCIPLES (apply to ANY label, not just examples)
════════════════════════════════════════
Two labels MUST be merged if they satisfy ANY of these:

  P1 — Root-word equivalence
      Any noun/verb/adjective/gerund form of the same root = same concept.
      "Teacher absenteeism" = "Absent teachers" = "Teachers not attending" = "Irregular teacher attendance"

  P2 — Synonym / paraphrase equivalence
      Replacing any word with a synonym leaves the meaning unchanged = same concept.
      "Low parental value for education" = "Parents devaluing education" = "Parental indifference to education"
      = "Parents not prioritizing schooling" = "Lack of parental interest in children's education"

  P3 — Cause / effect / barrier framing of the same phenomenon
      "X preventing Y" = "Y due to X" = "Lack of X" = "No X" = "X as a barrier"
      "Poverty preventing education" = "No money for school" = "Economic hardship blocking attendance"

  P4 — Subject-emphasis variants of the same action
      Shifting who is described (child vs. parent vs. school) for the same action = same concept.
      "Children doing domestic chores" = "Domestic chores keeping children home" = "Girls doing housework"

  P5 — Qualifier variants that don't change the core issue
      Adding/removing "frequent", "irregular", "low", "poor", "lack of", "limited", "inadequate" alone
      does not create a new concept.
      "Irregular attendance" = "Low attendance" = "Poor school attendance" = "Frequent absenteeism"

  P6 — Specific instance vs. general form
      A specific elaboration of the same barrier = same concept.
      "School 5 km from village" = "School far from home" = "Long distance to school"

DO NOT MERGE — keep as separate concepts when:
  • They describe genuinely different root causes (even if related)
    "Poverty" ≠ "Child labour" (child labour is a consequence of poverty, not the same)
    "Teacher shortage" ≠ "Teacher quality" (different issues)
  • One is a sub-type or specific form of the other (preserve the detail!)
    "Elopement" ≠ "Child marriage" (elopement is one form, keep both)
    "Child labour in agriculture" ≠ "Child labour due to poverty" (keep specific types)
    "Dropping out after primary" ≠ "Dropping out after middle school" (different stages)
  • One is a cause and the other is a completely different consequence
    "Substance abuse" ≠ "Domestic violence" (vice ≠ crime)
  • They represent fundamentally different stakeholder actions
    "School distance" ≠ "Lack of transport" (geography ≠ service)

IMPORTANT THEME RULES:
    • Theme must be the FULL theme name (not a number).
    • Do NOT use a theme name as a concept label.

IMPORTANT: When in doubt, keep concepts SEPARATE. The report needs rich detail,
not over-simplified categories. It is better to have two similar items than to lose
important detail by merging things that are related but distinct.

════════════════════════════════════════
FORMAT FOR CANONICAL LABELS
════════════════════════════════════════
  - Concise noun phrase, 3-8 words, no trailing punctuation
  - EXACT identical string for every item in the same group
  - Choose the most specific, descriptive label from the group

INPUT LIST:
{json.dumps(batch)}

════════════════════════════════════════
METHOD (follow in order)
════════════════════════════════════════
  1. Read ALL items.
  2. For each item, ask: "Is there another item that describes the SAME underlying issue?"
  3. Group variants of the same issue and assign ONE canonical label per group.
  4. SELF-CHECK: verify each merge — are these really the same issue, or related but distinct?
     If distinct, undo. But also check: did you miss any variants that should merge?
  5. Write output.

OUTPUT FORMAT:
Valid JSON object. Keys = input strings. Values = {{"concept": "...", "theme": "..."}}
RETURN ONLY VALID JSON. NO MARKDOWN. NO TEXT OUTSIDE JSON."""

        for attempt in range(1, MAX_RETRIES + 1):
            try:
                text = llm_provider.generate_text(
                    prompt,
                    max_tokens=REFINE_MAX_TOKENS,
                    temperature=0,
                    thinking_budget=REFINE_THINKING_BUDGET,
                )

                try:
                    batch_result = _parse_refinement_response(text, batch)
                except Exception:
                    batch_result = _repair_refinement_json_with_ai(text, batch, type_label)

                all_results.update(batch_result)
                save_refinement_checkpoint(checkpoint_path, all_results)
                print(f"      ✅ Batch {current_batch} saved ({len(batch_result)} items).")
                break

            except Exception as error:
                retryable = is_retryable_error(error)
                parse_error = isinstance(error, (json.JSONDecodeError, ValueError, TypeError))
                should_retry = (retryable or parse_error) and attempt < MAX_RETRIES
                print(f"      ⚠️ Batch {current_batch} attempt {attempt}/{MAX_RETRIES} failed: {error}")

                if should_retry:
                    delay = min(MAX_RETRY_SECONDS, BASE_RETRY_SECONDS * (2 ** (attempt - 1)))
                    delay += random.uniform(0, 0.5)
                    print(f"      🔁 Retrying in {delay:.1f}s...")
                    time.sleep(delay)
                    continue

                print(f"      ❌ Batch {current_batch} failed after {attempt} attempt(s).")
                break

    return all_results


def _cross_batch_dedup(all_results, type_label):
    """
    ══ OPTIMIZATION 5: Simplified cross-batch dedup ══
    Larger batch size (150 vs 80), fewer sub-batches needed.
    """
    if not llm_provider or not all_results:
        return all_results

    canonical_labels = list({v['concept'] for v in all_results.values()
                             if isinstance(v, dict) and v.get('concept')})
    if len(canonical_labels) <= 1:
        return all_results

    print(f"      🔁 Cross-batch dedup: {len(canonical_labels)} unique canonical labels — checking for duplicates...")

    full_label_map = {}
    total_merged = 0

    for batch_start in range(0, len(canonical_labels), CROSS_DEDUP_BATCH):
        batch_labels = canonical_labels[batch_start:batch_start + CROSS_DEDUP_BATCH]
        if len(batch_labels) <= 1:
            for lbl in batch_labels:
                full_label_map[lbl] = lbl
            continue

        batch_num = (batch_start // CROSS_DEDUP_BATCH) + 1
        total_batches = (len(canonical_labels) + CROSS_DEDUP_BATCH - 1) // CROSS_DEDUP_BATCH

        dedup_prompt = f"""You are a Data Cleaning Expert. The list below contains canonical {type_label} labels
produced by an AI working in SEPARATE batches. Because batches never saw each other, the SAME concept
frequently received DIFFERENT labels. Your job: identify all such duplicate labels and unify them.

════════════════════════════════════════
MERGE RULES
════════════════════════════════════════
MERGE any two labels that satisfy ANY of these:
  P1 — Root-word equivalence (noun/verb/adjective/gerund forms of same root)
  P2 — Synonym / paraphrase (replacing a word with its synonym keeps meaning)
  P3 — Cause/effect/barrier framing of the same phenomenon
  P4 — Subject-emphasis shift for the same action
  P5 — Qualifier variants ("irregular"/"low"/"poor" alone don't create new concepts)
  P6 — Specific elaboration of the same general barrier

DO NOT MERGE if root issues genuinely differ:
  "Poverty" ≠ "Child labour" | "Teacher shortage" ≠ "Teacher quality"

INPUT LABELS:
{json.dumps(batch_labels)}

OUTPUT: JSON object mapping EVERY input label to its final canonical label.
  - Unique labels map to themselves.
  - All labels in the same group map to the EXACT same string.
RETURN ONLY VALID JSON. NO MARKDOWN."""

        for attempt in range(1, MAX_RETRIES + 1):
            try:
                text = llm_provider.generate_text(
                    dedup_prompt,
                    max_tokens=REFINE_MAX_TOKENS,
                    temperature=0,
                    thinking_budget=REFINE_THINKING_BUDGET,
                )
                label_map = json.loads(_extract_json_block(text))
                if not isinstance(label_map, dict):
                    raise ValueError("Expected dict")

                for lbl in batch_labels:
                    full_label_map[lbl] = label_map.get(lbl, lbl)

                batch_merged = sum(1 for lbl in batch_labels if label_map.get(lbl, lbl) != lbl)
                total_merged += batch_merged
                if total_batches > 1:
                    print(f"         Batch {batch_num}/{total_batches}: merged {batch_merged} duplicate(s).")
                break

            except Exception as e:
                if attempt < MAX_RETRIES:
                    time.sleep(BASE_RETRY_SECONDS * (2 ** (attempt - 1)))
                    continue
                print(f"      ⚠️ Cross-batch dedup batch {batch_num} failed: {e}. Skipping batch.")
                for lbl in batch_labels:
                    full_label_map.setdefault(lbl, lbl)
                break

    if full_label_map:
        updated = {}
        for original, data in all_results.items():
            if not isinstance(data, dict):
                updated[original] = data
                continue
            old_concept = data.get('concept', '')
            new_concept = full_label_map.get(old_concept, old_concept)
            updated[original] = {**data, 'concept': new_concept}

        print(f"      ✅ Cross-batch dedup: merged {total_merged} duplicate label(s) total.")
        return updated

    return all_results


def _unified_semantic_dedup(df, type_label):
    """
    ══ OPTIMIZATION 2: Unified semantic dedup ══
    Single-pass deduplication across ALL themes simultaneously.
    Replaces BOTH per-theme semantic dedup AND cross-theme AI clustering.
    Typical call count: 1-2 per type (was ~20 combined).
    """
    if not llm_provider:
        return df

    MAX_ROUNDS = 2

    for round_num in range(1, MAX_ROUNDS + 1):
        # Collect all unique labels with dominant theme and total count
        concept_stats = df.groupby('Merged_Concept').agg(
            total=('Merged_Concept', 'size'),
            dominant_theme=('Theme', lambda x: x.value_counts().idxmax())
        ).reset_index()

        if len(concept_stats) <= 2:
            print(f"   ✅ Unified semantic dedup ({type_label}): only {len(concept_stats)} concepts — skipping.")
            break

        concept_stats = concept_stats.sort_values('total', ascending=False)
        label_list = concept_stats[['Merged_Concept', 'total', 'dominant_theme']].values.tolist()

        full_map = {}
        round_merges = 0

        for batch_start in range(0, len(label_list), UNIFIED_DEDUP_BATCH_SIZE):
            batch = label_list[batch_start:batch_start + UNIFIED_DEDUP_BATCH_SIZE]
            if len(batch) <= 1:
                for concept, _, _ in batch:
                    full_map[concept] = concept
                continue

            lines = [f"[{theme}] {concept}  ({total} mentions)" for concept, total, theme in batch]
            labels_text = "\n".join(lines)
            batch_labels = [concept for concept, _, _ in batch]

            prompt = f"""You are deduplicating {type_label.lower()} concept labels from an education report.
Labels are grouped by theme. Duplicates may exist WITHIN the same theme OR ACROSS different themes.

GOAL: Find labels that are linguistic duplicates — the SAME concept described with different words.

MERGE ONLY when:
• Two labels are paraphrases or rewordings of the EXACT SAME concept
  e.g. "Lack of teachers" = "Teacher shortage" = "Insufficient teachers"
  e.g. "Parents not valuing education" = "Low parental value for education"

KEEP SEPARATE — do NOT merge:
• A specific sub-type and its parent category
  "Child labour" ≠ "Poverty" (child labour is a consequence of poverty, not the same thing)
  "Elopement" ≠ "Child marriage" (elopement is one form, not the same as child marriage)
  "Weather barriers" ≠ "School distance" (different barriers)
  "Domestic violence" ≠ "Substance abuse" (related but different)
• Concepts with different root causes, even if related
• Concepts that need different interventions

IMPORTANT: When in doubt, keep concepts SEPARATE. It is better to have
two similar items than to lose important detail by over-merging.

When merging, use the label with MORE mentions as the canonical label.

INPUT LABELS (format: [Theme] Label (count)):
{labels_text}

OUTPUT: JSON object mapping EVERY input label (without theme prefix or mention counts) to its canonical label.
Only true duplicates share a canonical label. Everything else maps to itself.
RETURN ONLY VALID JSON. NO MARKDOWN."""

            for attempt in range(1, MAX_RETRIES + 1):
                try:
                    response = llm_provider.generate_text(
                        prompt,
                        max_tokens=REFINE_MAX_TOKENS,
                        temperature=0,
                        thinking_budget=REFINE_THINKING_BUDGET,
                    )
                    parsed = json.loads(_extract_json_block(response))
                    if not isinstance(parsed, dict):
                        raise ValueError("Expected dict")

                    known_labels = set(batch_labels)
                    for lbl in batch_labels:
                        mapped = parsed.get(lbl, lbl)
                        if mapped in known_labels:
                            full_map[lbl] = mapped
                        else:
                            full_map[lbl] = lbl

                    batch_merges = sum(1 for lbl in batch_labels if full_map.get(lbl, lbl) != lbl)
                    round_merges += batch_merges
                    break

                except Exception as e:
                    if attempt < MAX_RETRIES:
                        time.sleep(BASE_RETRY_SECONDS * (2 ** (attempt - 1)))
                        continue
                    print(f"      ⚠️ Unified dedup failed: {e}. Skipping batch.")
                    for lbl in batch_labels:
                        full_map.setdefault(lbl, lbl)
                    break

        if round_merges:
            print(f"   🔗 Unified semantic dedup ({type_label}) round {round_num}: {round_merges} label(s) merged.")
            df['Merged_Concept'] = df['Merged_Concept'].map(full_map).fillna(df['Merged_Concept'])
        else:
            print(f"   ✅ Unified semantic dedup ({type_label}) round {round_num}: converged — no more merges.")
            break

    return df


def _consolidate_cross_theme_concepts(df, type_label):
    """
    ══ OPTIMIZATION 4: Phase 1 only — exact-match, ZERO API calls ══
    Same label in 2+ themes → reassign all rows to dominant theme.
    AI-based Phase 2 is eliminated (handled by unified semantic dedup above).
    """
    concept_theme_counts = (
        df.groupby(['Merged_Concept', 'Theme'])
        .size()
        .reset_index(name='count')
    )

    idx_max = concept_theme_counts.groupby('Merged_Concept')['count'].idxmax()
    best_theme = concept_theme_counts.loc[idx_max]
    concept_to_dominant_theme = dict(zip(best_theme['Merged_Concept'], best_theme['Theme']))

    theme_counts_per_concept = concept_theme_counts.groupby('Merged_Concept')['Theme'].nunique()
    multi_theme_concepts = theme_counts_per_concept[theme_counts_per_concept > 1].index.tolist()

    if multi_theme_concepts:
        exact_merges = 0
        for concept in multi_theme_concepts:
            dominant = concept_to_dominant_theme[concept]
            mask = df['Merged_Concept'] == concept
            changed = (df.loc[mask, 'Theme'] != dominant).sum()
            if changed > 0:
                df.loc[mask, 'Theme'] = dominant
                exact_merges += changed

        if exact_merges:
            print(f"   🔄 Cross-theme consolidation ({type_label}): {len(multi_theme_concepts)} concept(s) "
                  f"in multiple themes — {exact_merges} row(s) reassigned to dominant theme.")

    return df


# ═══════════════════════════════════════════
# REPORT GENERATION
# ═══════════════════════════════════════════

def generate_report():
    print("🚀 Starting Optimized Final Report Generation Engine...")

    # 1. LOAD DATASETS
    print("   📂 Loading datasets...")
    try:
        df_raw = pd.read_csv('cleaned_data.csv')
        chal_exploded = pd.read_csv('exploded_challenges.csv')
        sol_exploded = pd.read_csv('exploded_solutions.csv')
        chal_map = pd.read_csv('challenge_mapping.csv')
        sol_map = pd.read_csv('solution_mapping.csv')
    except Exception as e:
        print(f"❌ Error: Required CSV files missing. {e}")
        return

    # Normalize mappings
    print("   ⚙️  Processing data and applying categories...")
    chal_map['Theme'] = chal_map['Theme'].apply(clean_theme_name)
    sol_map['Theme'] = sol_map['Theme'].apply(clean_theme_name)

    # CREATE df_c (Challenges) and apply environment logic
    df_c = chal_exploded.merge(chal_map, left_on='Challenges', right_on='Original', how='left')
    df_c['Theme'] = df_c['Theme'].fillna("Other Factors").astype(str)
    df_c['Theme'] = df_c['Theme'].apply(clean_theme_name)
    df_c['Environment'] = df_c['Challenges'].apply(categorize_environment_aggressive)

    df_chal_mapped = df_c

    # CREATE df_s (Solutions) and apply agency logic
    df_s = sol_exploded.merge(sol_map, left_on='Solutions', right_on='Original', how='left')
    df_s['Theme'] = df_s['Theme'].fillna("Other Factors").astype(str)
    df_s['Theme'] = df_s['Theme'].apply(clean_theme_name)
    df_s['Agency'] = df_s['Solutions'].apply(categorize_agency)

    # ═══════════════════════════════════════════
    # OPTIMIZED AI REFINEMENT PIPELINE
    # ═══════════════════════════════════════════
    #   Step 1: Refine concepts (larger batches → fewer calls)
    #   Step 2: Cross-batch dedup (simplified, larger batch)
    #   Step 3: Lexical consolidation (no API)
    #   Step 4: Unified semantic dedup (replaces per-theme + cross-theme AI)
    #   Step 5: Cross-theme exact-match reassignment (no API)

    top_chal = df_c['Merged_Concept'].value_counts().index.tolist()
    top_sol = df_s['Merged_Concept'].value_counts().index.tolist()

    # --- Refine Challenges ---
    chal_updates = refine_concepts_with_ai(top_chal, "Challenge")
    chal_updates = _cross_batch_dedup(chal_updates, "Challenge")
    if chal_updates:
        for old, new_data in chal_updates.items():
            mask = df_c['Merged_Concept'] == old
            df_c.loc[mask, 'Merged_Concept'] = new_data['concept']
            df_c.loc[mask, 'Theme'] = new_data['theme']

    # Lexical consolidation (no API calls)
    def _consolidate_canonical_labels(df_slice):
        label_counts = df_slice['Merged_Concept'].value_counts().to_dict()
        labels = list(label_counts.keys())
        parent = {lbl: lbl for lbl in labels}

        def _find(x):
            while parent[x] != x:
                parent[x] = parent[parent[x]]
                x = parent[x]
            return x

        def _union(a, b):
            ra, rb = _find(a), _find(b)
            if ra == rb:
                return
            if label_counts.get(ra, 0) >= label_counts.get(rb, 0):
                parent[rb] = ra
            else:
                parent[ra] = rb

        for i, la in enumerate(labels):
            ka = normalize_concept_key(la)
            for lb in labels[i + 1:]:
                kb = normalize_concept_key(lb)
                if _are_concepts_similar(ka, kb):
                    _union(la, lb)

        mapping = {lbl: _find(lbl) for lbl in labels}
        result = df_slice.copy()
        mapped = result['Merged_Concept'].map(mapping)
        result['Merged_Concept'] = mapped.where(mapped.notna(), result['Merged_Concept'])
        result = result.infer_objects(copy=False)
        return result

    def _consolidate_by_theme(df):
        frames = [_consolidate_canonical_labels(grp) for _, grp in df.groupby('Theme', sort=False)]
        return pd.concat(frames, ignore_index=True) if frames else df

    df_c = _consolidate_by_theme(df_c)

    # --- Refine Solutions ---
    sol_updates = refine_concepts_with_ai(top_sol, "Solution")
    sol_updates = _cross_batch_dedup(sol_updates, "Solution")
    if sol_updates:
        for old, new_data in sol_updates.items():
            mask = df_s['Merged_Concept'] == old
            df_s.loc[mask, 'Merged_Concept'] = new_data['concept']
            df_s.loc[mask, 'Theme'] = new_data['theme']

    df_s = _consolidate_by_theme(df_s)

    # --- Unified Semantic Dedup (replaces per-theme + cross-theme AI) ---
    print("   🔗 Running unified semantic deduplication for Challenges...")
    df_c = _unified_semantic_dedup(df_c, "Challenge")
    print("   🔗 Running unified semantic deduplication for Solutions...")
    df_s = _unified_semantic_dedup(df_s, "Solution")

    # --- Cross-theme exact-match reassignment (Phase 1 only, no API) ---
    print("   🔄 Running cross-theme concept consolidation (exact-match only)...")
    df_c = _consolidate_cross_theme_concepts(df_c, "Challenge")
    df_s = _consolidate_cross_theme_concepts(df_s, "Solution")

    # Re-clean themes
    df_c['Theme'] = df_c['Theme'].apply(clean_theme_name)
    df_s['Theme'] = df_s['Theme'].apply(clean_theme_name)

    df_chal_mapped = df_c

    # --- BASELINE METRIC CALCULATIONS ---
    TOTAL_CH_STATE = len(df_raw)
    for col in ['Participant Count', 'Men', 'Women', 'Children']:
        df_raw[col] = pd.to_numeric(df_raw[col], errors='coerce').fillna(0)

    TOTAL_PART_STATE = int(df_raw['Participant Count'].sum())
    NUM_CHAL_STATEMENTS = len(chal_exploded)
    NUM_SOL_STATEMENTS = len(sol_exploded)
    NUM_CHAL = NUM_CHAL_STATEMENTS
    NUM_SOL = NUM_SOL_STATEMENTS

    SOL_RATIO = (NUM_SOL / NUM_CHAL) if NUM_CHAL > 0 else 0

    m_total = int(df_raw['Men'].sum())
    w_total = int(df_raw['Women'].sum())
    c_total = int(df_raw['Children'].sum())
    df_raw['Others'] = df_raw['Participant Count'] - (df_raw['Men'] + df_raw['Women'] + df_raw['Children'])
    df_raw['Others'] = df_raw['Others'].clip(lower=0)
    o_total = int(df_raw['Others'].sum())

    w_perc = (w_total / TOTAL_PART_STATE * 100) if TOTAL_PART_STATE > 0 else 0
    m_perc = (m_total / TOTAL_PART_STATE * 100) if TOTAL_PART_STATE > 0 else 0
    c_perc = (c_total / TOTAL_PART_STATE * 100) if TOTAL_PART_STATE > 0 else 0
    o_perc = (o_total / TOTAL_PART_STATE * 100) if TOTAL_PART_STATE > 0 else 0

    avg_per_chaupal = TOTAL_PART_STATE / TOTAL_CH_STATE if TOTAL_CH_STATE > 0 else 0
    num_districts = df_raw['District'].nunique()
    num_themes = df_c['Theme'].nunique()

    theme_counts = df_c['Theme'].value_counts()
    top_3_themes = theme_counts.head(3)
    top_3_perc = (top_3_themes.sum() / NUM_CHAL * 100) if NUM_CHAL > 0 else 0

    agency_counts = df_s['Agency'].value_counts()
    ind_led = agency_counts.get('Individual-led', 0)
    comm_led = agency_counts.get('Community-led', 0)
    inst_led = agency_counts.get('Institutional', 0)

    ind_perc = (ind_led / NUM_SOL * 100) if NUM_SOL > 0 else 0
    comm_perc = (comm_led / NUM_SOL * 100) if NUM_SOL > 0 else 0
    inst_perc = (inst_led / NUM_SOL * 100) if NUM_SOL > 0 else 0

    comm_driven_perc = ind_perc + comm_perc

    doc = Document()

    # ═══════════════════════════════════════════
    # SECTION 1: EXECUTIVE SUMMARY
    # ═══════════════════════════════════════════
    print("   📝 Generating Section 1: Executive Summary...")
    doc.add_heading('1. EXECUTIVE SUMMARY', level=1)

    intro_p = doc.add_paragraph()
    intro_p.add_run(f"This comprehensive report presents an in-depth analysis of {TOTAL_CH_STATE:,} Shiksha Chaupal community dialogues conducted across Bihar, representing the collective voices of {TOTAL_PART_STATE:,} community members. These dialogues constitute one of the most extensive participatory consultations on education challenges in India, providing rich insights into grassroots barriers to education and community-driven solutions. The analysis encompasses {NUM_CHAL:,} individual challenges and {NUM_SOL:,} solutions, systematically categorized into {num_themes} primary thematic areas for comprehensive understanding.")

    doc.add_heading(f'🔑 KEY INSIGHT: Solution Coverage Ratio', level=2)
    p_ratio = doc.add_paragraph()
    run_ratio = p_ratio.add_run(f"Solution-to-Challenge Ratio: {SOL_RATIO:.2f}")
    run_ratio.bold = True

    if SOL_RATIO >= 1.0:
        ratio_text = f"This remarkable ratio demonstrates that communities identified {NUM_SOL:,} solutions for {NUM_CHAL:,} challenges. This transcends traditional deficit-based consultations where communities merely list problems. Instead, it reveals communities as active problem-solvers who think constructively about actionable interventions. This represents a paradigm shift in community engagement from problem identification to solution co-creation."
    elif SOL_RATIO >= 0.5:
        ratio_text = f"With {NUM_SOL:,} solutions proposed for {NUM_CHAL:,} challenges, communities are actively engaging in problem-solving. This indicates a constructive approach where participants are moving beyond just listing problems to identifying potential interventions."
    else:
        ratio_text = f"Communities identified {NUM_SOL:,} solutions alongside {NUM_CHAL:,} challenges. While the focus remains on highlighting barriers, there is an emerging capacity for solution-finding that can be further nurtured."

    doc.add_paragraph(ratio_text)

    doc.add_heading('Scale of Community Participation', level=2)
    scale_p = doc.add_paragraph()
    scale_p.add_run(f"• Geographic Reach: {TOTAL_CH_STATE:,} community dialogues conducted across {num_districts} districts in Bihar\n")
    scale_p.add_run(f"• Total Participants: {TOTAL_PART_STATE:,} community members actively engaged\n")
    scale_p.add_run(f"• Average Engagement: {avg_per_chaupal:.1f} participants per Chaupal, indicating strong community mobilization\n")
    scale_p.add_run(f"• Gender Representation: Women {w_perc:.1f}%, Children {c_perc:.1f}%, Men {m_perc:.1f}%, Others {o_perc:.1f}%\n")

    if w_perc > 50:
        gender_text = f"• Dominant female participation ({w_perc:.1f}%) signals authentic grassroots engagement rather than tokenistic consultation, as women are primary stakeholders in children's education"
    elif w_perc > m_perc:
        gender_text = f"• Strong female participation ({w_perc:.1f}%) highlights women's active role in discussing education challenges, outnumbering male participants ({m_perc:.1f}%)."
    else:
        gender_text = f"• The dialogues included diverse participation, with women contributing {w_perc:.1f}% of the voices, ensuring maternal perspectives are included."

    scale_p.add_run(gender_text)

    doc.add_heading('Dominant Challenge Themes', level=2)
    doc.add_paragraph(f"The thematic analysis reveals systemic patterns in education barriers. The top 3 themes collectively account for {top_3_perc:.1f}% of all challenges, indicating concentrated problem areas requiring prioritized intervention:")

    for theme, count in top_3_themes.items():
        t_perc = (count / NUM_CHAL * 100)
        s_count = len(df_s[df_s['Theme'] == theme])
        doc.add_paragraph(f"{theme}: {t_perc:.1f}% ({count:,} challenges, {s_count:,} solutions)", style='List Bullet')

    doc.add_heading('Community-Led vs System-Dependent Solutions', level=2)
    doc.add_paragraph("Solution agency analysis reveals community ownership patterns. The distribution demonstrates where communities see themselves as agents of change versus where they require external institutional support:")

    agency_p = doc.add_paragraph()
    agency_p.add_run(f"• Individual-led Solutions: {ind_perc:.1f}% ({ind_led:,} solutions) - Family-level actions including parental engagement, behavioral change, and household resource allocation\n")
    agency_p.add_run(f"• Community-led Solutions: {comm_perc:.1f}% ({comm_led:,} solutions) - Collective action including social mobilization, peer support networks, and community organizing\n")
    agency_p.add_run(f"• Institutional Solutions: {inst_perc:.1f}% ({inst_led:,} solutions) - Systemic interventions requiring government or CSO support including infrastructure, policy changes, and resource provision\n")

    if comm_driven_perc > 50:
        insight_text = f"• Critical Insight: {comm_driven_perc:.1f}% of solutions are community-driven (individual + community-led), demonstrating extraordinary grassroots capacity that partnerships must amplify rather than replace"
    else:
        insight_text = f"• Critical Insight: While {comm_driven_perc:.1f}% of solutions are community-driven, a significant portion ({inst_perc:.1f}%) requires institutional support, highlighting the need for strong government-community collaboration."

    crit_run = agency_p.add_run(insight_text)
    crit_run.bold = True

    doc.add_heading('🤝 Strategic Partnership Opportunity', level=2)

    if comm_driven_perc > 50:
        strat_text = f"The {comm_driven_perc:.1f}% proportion of community-led and individual-led solutions reveals extraordinary community ownership and problem-solving capacity. Strategic partnerships should operate on a community-strengthening model rather than community-replacing model. This means: (1) Amplifying existing community initiatives through capacity building and resource support, (2) Providing targeted institutional interventions ({inst_perc:.1f}%) for infrastructure, teacher capacity, and documentation systems that communities cannot address independently, (3) Facilitating community-to-community learning and peer exchange, (4) Advocating for policy changes that enable community solutions to scale. The partnership must recognize communities as co-creators and primary implementers, not merely beneficiaries."
    else:
        strat_text = f"With {inst_perc:.1f}% of solutions requiring institutional intervention, a collaborative partnership model is essential. This involves: (1) Government and CSOs addressing structural barriers like infrastructure and teacher shortages, (2) Strengthening the {comm_driven_perc:.1f}% of community-led initiatives to ensure sustainability, (3) Creating feedback loops where community needs directly inform policy implementation."

    doc.add_paragraph(strat_text)

    doc.add_page_break()

    # ═══════════════════════════════════════════
    # SECTION 2: GENERAL PARTICIPATION OVERVIEW
    # ═══════════════════════════════════════════
    print("   📝 Generating Section 2: Participation Overview...")
    doc.add_heading('2. GENERAL PARTICIPATION OVERVIEW', level=1)

    intro_para = doc.add_paragraph()
    run = intro_para.add_run("This section provides comprehensive analysis of participation patterns across geographic and demographic dimensions.")
    run.italic = True

    doc.add_heading('TABLE 1: Overall Participation Metrics', level=2)
    table1 = doc.add_table(rows=1, cols=3)
    table1.style = 'Table Grid'
    hdr_cells = table1.rows[0].cells
    for i, h in enumerate(['Metric', 'Count', 'Percentage (%)']):
        hdr_cells[i].text = h
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        set_cell_background(hdr_cells[i], "D9D9D9")

    def add_row_v3(t, metric, count, p_type):
        r = t.add_row().cells
        r[0].text = metric
        r[1].text = f"{int(count):,}" if isinstance(count, (int, float)) else str(count)
        if p_type == "none": r[2].text = "-"
        elif p_type == "full": r[2].text = "100%"
        else:
            p = (count / TOTAL_PART_STATE * 100) if TOTAL_PART_STATE > 0 else 0
            r[2].text = f"{p:.1f}%"

    add_row_v3(table1, "Total Number of Chaupals", TOTAL_CH_STATE, "none")
    add_row_v3(table1, "Total Number of Participants", TOTAL_PART_STATE, "full")
    add_row_v3(table1, "Average Participants per Chaupal", f"{avg_per_chaupal:.1f}", "none")
    add_row_v3(table1, "Men Participants", m_total, "calc")
    add_row_v3(table1, "Women Participants", w_total, "calc")
    add_row_v3(table1, "Children Participants", c_total, "calc")
    add_row_v3(table1, "Others (Unspecified)", o_total, "calc")

    doc.add_heading('2.3 District-wise Distribution of Reported Chaupals', level=2)
    dist_stats = df_raw.groupby('District').agg(Ch_Count=('id', 'count'), Part_Sum=('Participant Count', 'sum')).reset_index()
    dist_stats['Ch_Perc'] = (dist_stats['Ch_Count'] / TOTAL_CH_STATE) * 100
    dist_stats = dist_stats.sort_values('Ch_Count', ascending=False)

    table_dist = doc.add_table(rows=1, cols=5)
    table_dist.style = 'Table Grid'
    d_hdr = table_dist.rows[0].cells
    headers = ['District', 'Chaupals (N)', 'Chaupal %', 'Participants (N)', 'Participant %']
    for i, h in enumerate(headers):
        d_hdr[i].text = h
        d_hdr[i].paragraphs[0].runs[0].bold = True
        set_cell_background(d_hdr[i], "F2F2F2")

    for _, row in dist_stats.iterrows():
        r = table_dist.add_row().cells
        r[0].text, r[1].text = str(row['District']), f"{int(row['Ch_Count']):,}"
        r[2].text = f"{row['Ch_Perc']:.1f}%"
        r[3].text, r[4].text = f"{int(row['Part_Sum']):,}", f"{(row['Part_Sum']/TOTAL_PART_STATE)*100:.1f}%"

    top_3 = dist_stats.head(3)
    top_3_sum = top_3['Ch_Perc'].sum()
    dist_list_str = ", ".join([f"{row['District']} ({int(row['Ch_Count']):,}, {row['Ch_Perc']:.1f}%)" for _, row in top_3.iterrows()])

    doc.add_heading('Geographic Distribution Analysis', level=3)
    doc.add_paragraph(f"Geographic distribution shows concentration in {dist_list_str}. Together, these top 3 districts account for {top_3_sum:.1f}% of dialogues.")

    doc.add_heading('Demographic Composition Analysis', level=3)
    w_perc, m_perc, c_perc = (w_total/TOTAL_PART_STATE)*100, (m_total/TOTAL_PART_STATE)*100, (c_total/TOTAL_PART_STATE)*100
    ratio = w_perc / m_perc if m_perc > 0 else 0
    doc.add_paragraph(f"Women constitute {w_perc:.1f}% of participants, which is {( 'more than triple' if ratio >= 3 else 'significantly higher than' )} male participation ({m_perc:.1f}%).")

    doc.add_page_break()

    # ═══════════════════════════════════════════
    # SECTION 3: CORE CONTENT ANALYSIS
    # ═══════════════════════════════════════════
    print("   📝 Generating Section 3: Core Content Analysis...")
    doc.add_heading('3. CORE CONTENT ANALYSIS', level=1)

    unique_chal_count = chal_map['Merged_Concept'].nunique()
    unique_sol_count = sol_map['Merged_Concept'].nunique()
    chal_reduction = ((NUM_CHAL - unique_chal_count) / NUM_CHAL * 100) if NUM_CHAL > 0 else 0
    sol_reduction = ((NUM_SOL - unique_sol_count) / NUM_SOL * 100) if NUM_SOL > 0 else 0

    doc.add_paragraph(f"This section analyzes the substance of community dialogues - the challenges identified and solutions proposed. Communities articulated {NUM_CHAL:,} individual challenges and {NUM_SOL:,} individual solutions.")

    doc.add_heading('Overall Challenge & Solution Metrics', level=2)
    summary_table = doc.add_table(rows=1, cols=2)
    summary_table.style = 'Table Grid'
    hdr = summary_table.rows[0].cells
    hdr[0].text, hdr[1].text = 'Metric', 'Count'
    for cell in hdr:
        set_cell_background(cell, "D9D9D9")
        cell.paragraphs[0].runs[0].bold = True

    data_rows = [
        ("Total Challenges", f"{NUM_CHAL:,}"),
        ("Total Solutions", f"{NUM_SOL:,}"),
        ("Unique Challenges (after deduplication)", f"{unique_chal_count:,}"),
        ("Unique Solutions (after deduplication)", f"{unique_sol_count:,}"),
        ("Overall Solution-to-Challenge Ratio", f"{SOL_RATIO:.2f}"),
    ]
    for metric, val in data_rows:
        r = summary_table.add_row().cells
        r[0].text, r[1].text = metric, str(val)

    doc.add_heading('District-wise Engagement Depth', level=2)
    dist_chal_counts = chal_exploded.groupby('District').size().reset_index(name='C_Count')
    dist_sol_counts = sol_exploded.groupby('District').size().reset_index(name='S_Count')

    avg_df = dist_stats[['District', 'Ch_Count']].merge(dist_chal_counts, on='District').merge(dist_sol_counts, on='District')
    avg_df['Avg_C'] = avg_df['C_Count'] / avg_df['Ch_Count']
    avg_df['Avg_S'] = avg_df['S_Count'] / avg_df['Ch_Count']

    avg_table = doc.add_table(rows=1, cols=3)
    avg_table.style = 'Table Grid'
    a_hdr = avg_table.rows[0].cells
    a_hdr[0].text, a_hdr[1].text, a_hdr[2].text = 'District', 'Avg Challenges/Chaupal', 'Avg Solutions/Chaupal'
    for cell in a_hdr: set_cell_background(cell, "F2F2F2")

    for _, row in avg_df.sort_values('Avg_C', ascending=False).iterrows():
        r = avg_table.add_row().cells
        r[0].text, r[1].text, r[2].text = str(row['District']), f"{row['Avg_C']:.2f}", f"{row['Avg_S']:.2f}"

    doc.add_heading('Solution-to-Challenge Ratio: A Paradigm Shift', level=2)
    doc.add_paragraph(f"The overall solution-to-challenge ratio of {SOL_RATIO:.2f} represents a fundamental paradigm shift in community consultation methodology. Traditional deficit-based consultations focus solely on problem identification, treating communities as problem containers. The Shiksha Chaupal model demonstrates that when communities are engaged as problem-solvers rather than merely problem-identifiers, they actively think constructively about solutions. This {NUM_SOL:,} solutions for {NUM_CHAL:,} challenges ratio indicates that every articulated challenge was matched with actionable solution thinking, demonstrating community agency and constructive engagement. This asset-based approach recognizes communities as repositories of contextual knowledge and innovative problem-solving capacity.")

    doc.add_page_break()

    # ═══════════════════════════════════════════
    # SECTION 4: THEMATIC ANALYSIS (OPTIMIZED INSIGHTS)
    # ═══════════════════════════════════════════
    print("   📝 Generating Section 4: Thematic Analysis...")
    doc.add_heading('4. THEMATIC ANALYSIS', level=1)

    # TABLE: Challenge by Theme
    doc.add_heading('Overall Challenge Distribution by Theme', level=2)
    t_chal = doc.add_table(rows=1, cols=3); t_chal.style = 'Table Grid'
    h_chal = t_chal.rows[0].cells
    h_chal[0].text, h_chal[1].text, h_chal[2].text = 'Theme', 'Count', '%'
    for c in h_chal: set_cell_background(c, "D9D9D9")

    theme_counts = df_c['Theme'].value_counts()
    for theme, count in theme_counts.items():
        r = t_chal.add_row().cells
        r[0].text, r[1].text, r[2].text = theme, str(count), f"{(count/NUM_CHAL_STATEMENTS*100):.1f}%"

    # TABLE: Solution Agency
    doc.add_heading('Solution Distribution by Agency', level=2)
    t_agency = doc.add_table(rows=1, cols=3); t_agency.style = 'Table Grid'
    h_age = t_agency.rows[0].cells
    h_age[0].text, h_age[1].text, h_age[2].text = 'Agency Type', 'Count', '%'
    for c in h_age: set_cell_background(c, "F2F2F2")

    agency_counts = df_s['Agency'].value_counts()
    for agency, count in agency_counts.items():
        r = t_agency.add_row().cells
        r[0].text, r[1].text, r[2].text = agency, str(count), f"{(count/NUM_SOL_STATEMENTS*100):.1f}%"

    # TABLE: Challenge Environment
    doc.add_heading('Challenge Distribution by Environment', level=2)
    t_env = doc.add_table(rows=1, cols=3); t_env.style = 'Table Grid'
    h_env = t_env.rows[0].cells
    h_env[0].text, h_env[1].text, h_env[2].text = 'Environment', 'Count', '%'
    for c in h_env: set_cell_background(c, "D9D9D9")

    env_counts = df_c['Environment'].value_counts()
    for env, count in env_counts.items():
        r = t_env.add_row().cells
        r[0].text, r[1].text, r[2].text = env, str(count), f"{(count/NUM_CHAL_STATEMENTS*100):.1f}%"

    # --- INDIVIDUAL THEME DEEP-DIVES ---
    doc.add_page_break()

    themes_to_process = list(theme_counts.index[:10])
    if "Other Factors" in themes_to_process:
        themes_to_process.remove("Other Factors")
    if "Other Factors" in theme_counts.index:
        themes_to_process.append("Other Factors")

    for theme_idx, theme in enumerate(themes_to_process, 1):
        doc.add_heading(f'4.{theme_idx} {theme.upper()}', level=2)

        t_c = df_c[df_c['Theme'] == theme]
        t_s = df_s[df_s['Theme'] == theme]

        c_count = len(t_c)
        s_count = len(t_s)
        c_perc_theme = (c_count / NUM_CHAL * 100) if NUM_CHAL > 0 else 0
        sol_cov = (s_count / c_count) if c_count > 0 else 0
        u_c = t_c['Merged_Concept'].nunique()
        u_s = t_s['Merged_Concept'].nunique()

        m_para = doc.add_paragraph()
        m_para.add_run(f"Scale: {c_count:,} challenges ({c_perc_theme:.1f}% of total dataset) | {s_count:,} solutions\n").bold = True
        m_para.add_run(f"Solution Coverage: {sol_cov:.2f} solutions per challenge")

        # Challenge Landscape
        doc.add_heading('Challenge Landscape', level=4)
        if not t_c.empty:
            env_pref = t_c['Environment'].value_counts(normalize=True).idxmax()
            doc.add_paragraph(f"The landscape for '{theme}' is primarily localized within the {env_pref} environment. This suggests that interventions must be targeted at this level for maximum impact.")

        doc.add_heading("Top Recurring Challenges", level=5)

        # ══ OPTIMIZATION 3: Two-pass approach for batched insights ══
        # Pass 1: Determine which concepts to show (50% coverage + min 5 items)
        grouped_challenges = []
        challenges_with_key = assign_concept_groups(t_c, concept_column='Merged_Concept')

        for concept_key, group in challenges_with_key.groupby('Concept_Group'):
            mention_count = len(group)
            display_concept = group['Merged_Concept'].value_counts().idxmax()
            grouped_challenges.append((display_concept, mention_count, concept_key))

        grouped_challenges = sorted(grouped_challenges, key=lambda item: item[1], reverse=True)
        total_theme_chal = len(t_c)
        cumulative_count = 0

        concepts_to_show = []
        for concept, count, concept_key in grouped_challenges:
            cumulative_count += count
            coverage_perc = (cumulative_count / total_theme_chal) * 100
            item_perc = (count / total_theme_chal) * 100

            concept_group_rows = challenges_with_key[challenges_with_key['Concept_Group'] == concept_key]
            original_texts = concept_group_rows['Challenges'].tolist()

            # Find representative quote
            rep_quote = concept
            rep_district = "Unknown"
            if not concept_group_rows.empty:
                text_lengths = concept_group_rows['Challenges'].astype(str).str.len()
                rep_idx = text_lengths.idxmax()
                rep_row = concept_group_rows.loc[rep_idx]
                rep_quote = str(rep_row.get('Challenges', concept)).strip() or concept
                if 'District' in concept_group_rows.columns and pd.notna(rep_row.get('District')):
                    rep_district = str(rep_row.get('District')).strip() or "Unknown"

            # Compute metadata for insight generation
            districts = concept_group_rows['District'].nunique() if not concept_group_rows.empty and 'District' in concept_group_rows.columns else 0
            env_text = "Not available"
            if not concept_group_rows.empty and 'Environment' in concept_group_rows.columns:
                env_mix = concept_group_rows['Environment'].value_counts(normalize=True)
                if not env_mix.empty:
                    env_text = env_mix.index[0]

            concepts_to_show.append({
                'concept': concept,
                'count': count,
                'share': item_perc,
                'concept_key': concept_key,
                'concept_rows': concept_group_rows,
                'original_texts': original_texts,
                'rep_quote': rep_quote,
                'rep_district': rep_district,
                'districts': districts,
                'env_text': env_text,
            })

            if coverage_perc >= 50 and len(concepts_to_show) >= 5:
                break
            if len(concepts_to_show) >= 15:
                break

        # Pass 2: Batch-generate insights for all concepts in this theme
        insights_map = _batch_challenge_insights(theme, concepts_to_show, t_c, t_s)

        # Pass 3: Render each concept with pre-computed insights
        for i, data in enumerate(concepts_to_show, 1):
            concept = data['concept']
            count = data['count']
            item_perc = data['share']
            rep_quote = data['rep_quote']
            rep_district = data['rep_district']

            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(36)
            p.paragraph_format.first_line_indent = Pt(-18)
            p.add_run(f"{i}. {concept}").bold = True
            p.add_run(f" ({count} mentions, {item_perc:.1f}%)")

            # Get insights from batch results, or use fallback
            insight_paragraphs = insights_map.get(concept, None)
            if not insight_paragraphs:
                fallback = _challenge_item_insight_fallback(
                    theme, concept, item_perc, count,
                    data['districts'], data['env_text']
                )
                insight_paragraphs = [fallback]

            quote_length = len(rep_quote) if rep_quote else 0
            paragraphs_to_show = insight_paragraphs[:1] if quote_length < 100 else insight_paragraphs

            for insight_text in paragraphs_to_show:
                insight_para = doc.add_paragraph()
                insight_para.paragraph_format.left_indent = Pt(54)
                insight_para.add_run(insight_text)

            quote_para = doc.add_paragraph()
            quote_para.paragraph_format.left_indent = Pt(54)
            quote_para.add_run(f"Voice from the ground ({rep_district}): \"{rep_quote}\"").italic = True

        # Solution Ecosystem
        doc.add_heading('Solution Ecosystem', level=4)
        if not t_s.empty:
            total_theme_sol = len(t_s)
            agency_counts = t_s['Agency'].value_counts(normalize=True)
            agency_main = agency_counts.idxmax()
            agency_perc = agency_counts.max() * 100

            doc.add_paragraph(f"Communities proposed {total_theme_sol:,} solutions to address this theme. The solution ecosystem demonstrates {agency_main} agency with {agency_perc:.1f}% of solutions being {agency_main}.")

            doc.add_heading("Most Frequently Proposed Solutions", level=5)

            valid_solutions = t_s[t_s['Merged_Concept'].apply(is_valid_solution)].copy()
            valid_solutions = assign_concept_groups(valid_solutions, concept_column='Merged_Concept')

            grouped_solutions = []
            for concept_key, group in valid_solutions.groupby('Concept_Group'):
                mention_count = len(group)
                display_concept = group['Merged_Concept'].value_counts().idxmax()
                grouped_solutions.append((display_concept, mention_count, concept_key))

            top_solutions = sorted(grouped_solutions, key=lambda item: item[1], reverse=True)[:5]

            for rank, (concept, count, concept_key) in enumerate(top_solutions, 1):
                item_perc = (count / total_theme_sol) * 100
                original_texts = valid_solutions[valid_solutions['Concept_Group'] == concept_key]['Solutions'].tolist()
                rep_quote = max(original_texts, key=len) if original_texts else concept

                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(36)
                p.paragraph_format.first_line_indent = Pt(-18)
                p.add_run(f"{rank}. {concept}").bold = True
                p.add_run(f" ({count} mentions, {item_perc:.1f}%)")
                p.add_run(f"\n   Community Proposal: \"{rep_quote}\"").italic = True

    doc.add_page_break()

    # ═══════════════════════════════════════════
    # SECTION 5: DISTRICT PROFILES
    # ═══════════════════════════════════════════
    print("   📝 Generating Section 5: District Profiles...")
    doc.add_heading('5. DISTRICT PROFILES', level=1)

    doc.add_paragraph("This section provides comprehensive profiles for top-performing districts by Chaupal count, including quantitative metrics, thematic breakdowns, and top challenges/solutions specific to each district. These profiles enable district-specific intervention design and comparative analysis across geographies.")

    doc.add_heading('District Performance Overview Table', level=2)

    d_stats = df_raw.groupby('District').agg(
        Chaupals=('id', 'nunique'),
        Participants=('Participant Count', 'sum')
    ).reset_index()

    c_counts = chal_exploded.groupby('District').size().reset_index(name='Challenges')
    s_counts = sol_exploded.groupby('District').size().reset_index(name='Solutions')

    dist_overview = d_stats.merge(c_counts, on='District', how='left').merge(s_counts, on='District', how='left').fillna(0)
    dist_overview['Ratio'] = dist_overview['Solutions'] / dist_overview['Challenges']
    dist_overview = dist_overview.sort_values('Chaupals', ascending=False)

    table = doc.add_table(rows=1, cols=6)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    headers = ['District', 'Chaupals', 'Participants', 'Challenges', 'Solutions', 'Ratio']
    for i, h in enumerate(headers):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
        set_cell_background(hdr[i], "D9D9D9")

    for _, row in dist_overview.iterrows():
        r = table.add_row().cells
        r[0].text = str(row['District'])
        r[1].text = f"{int(row['Chaupals']):,}"
        r[2].text = f"{int(row['Participants']):,}"
        r[3].text = f"{int(row['Challenges']):,}"
        r[4].text = f"{int(row['Solutions']):,}"
        r[5].text = f"{row['Ratio']:.2f}"

    note_p = doc.add_paragraph()
    note_run = note_p.add_run("Note: Ratio = Solutions ÷ Challenges. Values >1.0 indicate more solutions than challenges identified.")
    note_run.italic = True
    note_run.font.size = Pt(9)

    dist_list = dist_overview['District'].tolist()
    if 'Others' in dist_list:
        dist_list.remove('Others')
        dist_list.append('Others')

    for i, dist in enumerate(dist_list, 1):
        doc.add_heading(f'5.{i} {dist.upper()}', level=2)
        d_raw = df_raw[df_raw['District'] == dist]

        d_chal = df_c[df_c['District'] == dist]
        d_sol = df_s[df_s['District'] == dist]

        dist_chaupals = len(d_raw)
        dist_participants = int(d_raw['Participant Count'].sum())
        chaupal_perc = (dist_chaupals / TOTAL_CH_STATE * 100) if TOTAL_CH_STATE > 0 else 0

        m_d = d_raw['Men'].sum()
        w_d = d_raw['Women'].sum()
        c_d = d_raw['Children'].sum()

        m_p = (m_d / dist_participants * 100) if dist_participants > 0 else 0
        w_p = (w_d / dist_participants * 100) if dist_participants > 0 else 0
        c_p = (c_d / dist_participants * 100) if dist_participants > 0 else 0

        dist_chal_count = len(d_chal)
        dist_sol_count = len(d_sol)
        dist_ratio = (dist_sol_count / dist_chal_count) if dist_chal_count > 0 else 0

        avg_c = dist_chal_count / dist_chaupals if dist_chaupals > 0 else 0
        avg_s = dist_sol_count / dist_chaupals if dist_chaupals > 0 else 0

        doc.add_heading('A. Quantitative Snapshot', level=3)
        snap_p = doc.add_paragraph()
        snap_p.add_run(f"Chaupals: {dist_chaupals:,} ({chaupal_perc:.1f}% of total)\n")
        snap_p.add_run(f"Total Participants: {dist_participants:,}\n")
        snap_p.add_run(f"Demographics: Men {m_p:.1f}%, Women {w_p:.1f}%, Children {c_p:.1f}%\n")
        snap_p.add_run(f"Challenges: {dist_chal_count:,} | Solutions: {dist_sol_count:,}\n")
        snap_p.add_run(f"Solution Efficiency: {dist_ratio:.2f}\n")
        snap_p.add_run(f"Average per Chaupal: {avg_c:.1f} challenges, {avg_s:.1f} solutions")

        if d_chal.empty:
            doc.add_paragraph("No challenge data available for this district.")
            continue

        total_dist_chal = len(d_chal)
        theme_counts = d_chal['Theme'].value_counts()

        doc.add_heading('Thematic Breakdown & Examples', level=3)

        for theme, count in theme_counts.items():
            perc = (count / total_dist_chal) * 100

            p_theme = doc.add_paragraph()
            p_theme.paragraph_format.space_before = Pt(6)
            run = p_theme.add_run(f"• {theme} ({perc:.1f}%)")
            run.bold = True

            theme_c_rows = d_chal[d_chal['Theme'] == theme]
            top_challenges = theme_c_rows['Merged_Concept'].value_counts().head(2).index.tolist()

            theme_s_rows = d_sol[d_sol['Theme'] == theme]

            if theme_s_rows.empty:
                top_solutions = []
            else:
                theme_s_rows = theme_s_rows.copy()
                theme_s_rows['Merged_Concept'] = theme_s_rows['Merged_Concept'].fillna("Uncategorized").astype(str)
                valid_mask = theme_s_rows['Merged_Concept'].apply(is_valid_solution)
                valid_s_rows = theme_s_rows[valid_mask]
                top_solutions = valid_s_rows['Merged_Concept'].value_counts().head(2).index.tolist()

            if top_challenges:
                p_c = doc.add_paragraph()
                p_c.paragraph_format.left_indent = Pt(18)
                p_c.add_run("Challenges: ").bold = True
                p_c.add_run("; ".join(top_challenges))

            if top_solutions:
                p_s = doc.add_paragraph()
                p_s.paragraph_format.left_indent = Pt(18)
                p_s.add_run("Solutions: ").bold = True
                p_s.add_run("; ".join(top_solutions))

    # ═══════════════════════════════════════════
    # SECTION 6: UNIQUE INSIGHTS
    # ═══════════════════════════════════════════
    print("   📝 Generating Section 6: Unique Insights...")
    doc.add_page_break()
    doc.add_heading('6. UNIQUE INSIGHTS', level=1)

    def get_unique_examples(agency_type, count=5):
        subset = df_s[df_s['Agency'] == agency_type].copy()
        if subset.empty: return []

        others = subset[subset['Theme'] == 'Other Factors']
        freq = subset['Merged_Concept'].value_counts()
        unique_concepts = freq[freq == 1].index.tolist()
        unique_rows = subset[subset['Merged_Concept'].isin(unique_concepts)]

        candidates = pd.concat([others, unique_rows]).drop_duplicates(subset=['Solutions'])
        candidates = candidates[candidates['Solutions'].str.len() > 20]
        candidates['len'] = candidates['Solutions'].str.len()
        top_candidates = candidates.sort_values('len', ascending=False).head(count)

        results = []
        for _, row in top_candidates.iterrows():
            dist = str(row['District'])
            sol = row['Solutions'].strip()
            location_str = f"{dist}"
            results.append(f"\"{sol}\"\n   📍 {location_str}")
        return results

    doc.add_heading('Individual-led Unique Solutions', level=2)
    doc.add_paragraph("Highlights of innovative or distinct solutions proposed by individuals that stand out from common themes:")
    for ex in get_unique_examples('Individual-led', 5):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(ex)

    doc.add_heading('Community-led Unique Solutions', level=2)
    doc.add_paragraph("Highlights of collective actions or community-driven innovations:")
    for ex in get_unique_examples('Community-led', 5):
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(ex)

    doc.add_heading('Expectations from Government & CSOs', level=2)
    doc.add_paragraph("Key expectations and demands expressed by the community for systemic support:")

    inst_subset = df_s[df_s['Agency'] == 'Institutional']
    if not inst_subset.empty:
        top_inst = inst_subset['Merged_Concept'].value_counts().head(5)
        for concept, count in top_inst.items():
            quotes = inst_subset[inst_subset['Merged_Concept'] == concept]['Solutions'].tolist()
            quotes = [q for q in quotes if isinstance(q, str) and len(q) > 10]
            if not quotes: continue

            best_quote = max(quotes, key=len)

            p = doc.add_paragraph(style='List Number')
            p.add_run(f"{concept}").bold = True
            p.add_run(f" ({count} mentions)")
            p.add_run(f"\n   Community Voice: \"{best_quote}\"").italic = True

    print("   💾 Saving document...")

    # ═══════════════════════════════════════════
    # SECTION 7: CONCLUSION
    # ═══════════════════════════════════════════
    print("   📝 Generating Section 7: Conclusion...")
    doc.add_page_break()
    doc.add_heading('7. CONCLUSION', level=1)

    doc.add_heading('🎯 TRANSFORMATIVE INSIGHT: Community as Solution Architects', level=2)
    doc.add_paragraph(f"This analysis of {TOTAL_CH_STATE:,} Shiksha Chaupals reveals a fundamental truth: rural communities are not passive recipients of development interventions but sophisticated problem-solvers capable of designing culturally-appropriate, sustainable solutions to education barriers. The {SOL_RATIO:.2f} solution coverage ratio demonstrates that communities are generating solutions at nearly the same rate as identifying challenges.")

    doc.add_heading('The Interconnected Challenge Reality', level=2)
    top_themes_list = top_3_themes.index.tolist()
    theme_str = ", ".join(top_themes_list)
    doc.add_paragraph(f"Challenges do not exist in isolation. The dominance of themes like {theme_str} suggests a complex interplay of factors. For instance, economic barriers often amplify documentation issues, while infrastructure gaps can drive families toward alternative schooling options. Understanding these interconnections is essential for designing effective interventions.")

    doc.add_heading('Community Agency Excellence', level=2)
    comm_sols = df_s[df_s['Agency'] == 'Community-led']['Merged_Concept'].value_counts().head(4).index.tolist()
    if comm_sols:
        comm_examples = ", ".join([s.lower() for s in comm_sols])
        doc.add_paragraph(f"Communities have demonstrated remarkable innovation through initiatives such as {comm_examples}. This agency must be recognized, celebrated, and supported—not replaced by external solutions.")
    else:
        doc.add_paragraph("Communities have demonstrated remarkable innovation through collective action and peer support mechanisms. This agency must be recognized, celebrated, and supported—not replaced by external solutions.")

    doc.add_heading('Systemic Support Imperative', level=2)
    inst_sols = df_s[df_s['Agency'] == 'Institutional']['Merged_Concept'].value_counts().head(4).index.tolist()
    if inst_sols:
        inst_examples = ", ".join([s.lower() for s in inst_sols])
        doc.add_paragraph(f"While community agency is exceptional, certain barriers require institutional action. Issues such as {inst_examples} cannot be resolved through community effort alone. The path forward requires strategic partnerships that amplify community strengths while providing systemic support.")
    else:
        doc.add_paragraph("While community agency is exceptional, certain barriers require institutional action. Infrastructure gaps, documentation bottlenecks, and resource shortages cannot be resolved through community effort alone. The path forward requires strategic partnerships that amplify community strengths while providing systemic support.")

    doc.add_heading('🤝 Strategic Partnership Opportunity', level=2)
    doc.add_paragraph("The optimal collaboration model combines community-led cultural change initiatives with institutional support for infrastructure and policy barriers. Government agencies, NGOs, and CSOs should position themselves as resource partners and accountability allies—not solution designers—enabling communities to scale their own innovations while addressing systemic gaps.")

    doc.save('Final_Shiksha_Report.docx')

    # Print optimization summary
    if llm_provider:
        print(f"\n   📊 Total API calls: {llm_provider.call_counter}")
    print("\n🏁 SUCCESS! Optimized report generated.")


if __name__ == "__main__":
    generate_report()
